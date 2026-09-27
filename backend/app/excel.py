"""Exportación a Excel (.xlsx) con openpyxl."""
from io import BytesIO
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlmodel import Session

from . import calc
from .db import get_settings

HEAD_FILL = PatternFill("solid", fgColor="1F4E46")
HEAD_FONT = Font(bold=True, color="FFFFFF")
DUR = "[h]:mm:ss"          # formato de duración que suma más de 24 h
DUR_NEG = "[h]:mm:ss;-[h]:mm:ss"


def _sheet(wb, title, headers, widths):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        c = ws.cell(row=1, column=i)
        c.fill, c.font, c.alignment = HEAD_FILL, HEAD_FONT, Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    return ws


def _d(seconds: float) -> float:
    """Segundos → fracción de día (valor de tiempo de Excel)."""
    return seconds / 86400


def build_xlsx(s: Session, desde: Optional[str] = None, hasta: Optional[str] = None) -> bytes:
    st = get_settings(s)
    wb = Workbook()
    wb.remove(wb.active)

    # 1) Jornadas
    ws = _sheet(wb, "Jornadas", ["ID", "Fecha", "Entrada", "Salida", "Duración", "Horas (decimal)", "Nota"],
                [6, 12, 10, 10, 12, 15, 40])
    shifts = calc.list_shifts(s, desde, hasta)
    for sh in shifts:
        ini = calc.to_local(sh.start, st)
        fin = calc.to_local(sh.end, st) if sh.end else None
        secs = calc.shift_seconds(sh)
        ws.append([sh.id, ini.date(), ini.time().replace(tzinfo=None),
                   fin.time().replace(tzinfo=None) if fin else "en curso",
                   _d(secs), round(secs / 3600, 2), sh.note])
        r = ws.max_row
        ws.cell(r, 2).number_format = "dd/mm/yyyy"
        ws.cell(r, 3).number_format = ws.cell(r, 4).number_format = "hh:mm"
        ws.cell(r, 5).number_format = DUR
    n = ws.max_row
    ws.append(["", "TOTAL", "", "", f"=SUM(E2:E{n})" if n > 1 else 0, f"=SUM(F2:F{n})" if n > 1 else 0, ""])
    ws.cell(ws.max_row, 5).number_format = DUR
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)

    # 2) Resumen diario
    ws = _sheet(wb, "Resumen diario", ["Fecha", "Trabajado", "Objetivo", "Extra", "Faltante", "Ajuste manual", "Nota ajuste"],
                [12, 12, 12, 12, 12, 14, 30])
    days = calc.daily_summary(s, desde, hasta)
    for d in days:
        ws.append([calc.date.fromisoformat(d["fecha"]), _d(d["trabajado_segundos"]), _d(st.daily_hours * 3600),
                   _d(d["extra_segundos"]), _d(d["faltante_segundos"]), "Sí" if d["ajuste_manual"] else "No",
                   d["nota_ajuste"]])
        r = ws.max_row
        ws.cell(r, 1).number_format = "dd/mm/yyyy"
        for col in (2, 3, 4):
            ws.cell(r, col).number_format = DUR
        ws.cell(r, 5).number_format = DUR_NEG
    n = ws.max_row
    ws.append(["TOTAL", f"=SUM(B2:B{n})" if n > 1 else 0, "", f"=SUM(D2:D{n})" if n > 1 else 0,
               f"=SUM(E2:E{n})" if n > 1 else 0, "", ""])
    for col in (2, 4):
        ws.cell(ws.max_row, col).number_format = DUR
    ws.cell(ws.max_row, 5).number_format = DUR_NEG
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)

    # 3) Banco de horas (saldo total, sin filtro de fechas)
    bank = calc.bank_status(s)
    ws = _sheet(wb, "Banco de horas", ["Fecha", "Tipo", "Tiempo", "Nota"], [12, 10, 12, 40])
    for m in bank["movimientos"]:
        ws.append([calc.date.fromisoformat(m["fecha"]), m["tipo"], _d(m["segundos"]), m["nota"]])
        ws.cell(ws.max_row, 1).number_format = "dd/mm/yyyy"
        ws.cell(ws.max_row, 3).number_format = DUR_NEG
    ws.append([])
    lab, rel = bank["desglose"]["laboral"], bank["desglose"]["reloj"]
    sg = "-" if lab["signo"] < 0 else ""
    for k, v in [("Extras acumuladas", bank["extras_acumuladas"]), ("Faltantes", bank["faltantes"]),
                 ("Horas usadas", bank["usadas"]), ("Ajustes a favor", bank["ajustes_a_favor"]),
                 ("SALDO (hh:mm:ss)", bank["saldo"]),
                 (f"Saldo laboral (jornada {st.daily_hours} h, mes {st.workdays_per_month} jornadas)",
                  f"{sg}{lab['meses']} meses, {lab['dias']} días, {lab['horas']} h, {lab['minutos']} min, {lab['segundos']} s"),
                 ("Saldo reloj (día 24 h, mes 30 días)",
                  f"{sg}{rel['meses']} meses, {rel['dias']} días, {rel['horas']} h, {rel['minutos']} min, {rel['segundos']} s")]:
        ws.append([k, "", v])
        ws.cell(ws.max_row, 1).font = Font(bold=True)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
