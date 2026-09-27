import json
import os
import tempfile

os.environ["DATABASE_URL"] = "sqlite:///" + os.path.join(tempfile.mkdtemp(), "t.db")
os.environ["OWNER_EMAIL"] = "dueno@example.com"

import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from io import BytesIO  # noqa: E402
from sqlmodel import Session, select  # noqa: E402

from backend.app import alarms, assistant, auth, calc, history, usage  # noqa: E402
from backend.app import google_integration as gi  # noqa: E402
from backend.app.db import (AiUsage, Alarm, BankMovement, BlockedEmail, ChatMessage, DayOverride, GoogleToken,  # noqa: E402
                            LoginSession, OAuthState, Settings, Shift, User, engine, init_db, utcnow)
from backend.app.main import app  # noqa: E402

OWNER = "dueno@example.com"
FRIEND = "amigo@example.com"


@pytest.fixture(autouse=True)
def clean_db():
    init_db()
    with Session(engine) as s:
        for model in (AiUsage, ChatMessage, Alarm, GoogleToken, DayOverride, BankMovement, Shift, Settings, LoginSession, OAuthState,
                      BlockedEmail, User):
            for row in s.exec(select(model)).all():
                s.delete(row)
        s.commit()
    yield


def make_user(email: str, name: str = "Test") -> tuple[int, str]:
    """Crea el usuario (si no existe) con una sesión iniciada. Devuelve (id, token de la cookie)."""
    with Session(engine) as s:
        u = s.exec(select(User).where(User.email == email)).first() or User(email=email, name=name)
        s.add(u)
        s.commit()
        s.refresh(u)
        return u.id, auth.create_session(s, u.id)


def client(email: str | None = None) -> TestClient:
    c = TestClient(app)
    if email:
        c.cookies.set(auth.COOKIE, make_user(email)[1])
    return c


def block(email: str) -> None:
    with Session(engine) as s:
        auth.block(s, email)


# ---------------------------------------------------------------- acceso
def test_requires_login():
    c = client()
    assert c.get("/api/health").status_code == 200
    assert c.get("/api/status").status_code == 401
    assert c.get("/api/shifts").status_code == 401
    assert c.post("/api/clock-in", json={}).status_code == 401
    assert c.get("/api/admin/usuarios").status_code == 401
    c.cookies.set(auth.COOKIE, "token-inventado")
    assert c.get("/api/status").status_code == 401


def test_session_token_is_stored_hashed_and_expires():
    uid, token = make_user(OWNER)
    with Session(engine) as s:
        rows = s.exec(select(LoginSession)).all()
        assert len(rows) == 1 and rows[0].token_hash != token and token not in rows[0].token_hash
    c = client()
    c.cookies.set(auth.COOKIE, token)
    assert c.get("/api/status").status_code == 200
    with Session(engine) as s:  # vence la sesión
        row = s.exec(select(LoginSession)).first()
        row.expires_at = utcnow() - auth.timedelta(seconds=1)
        s.add(row)
        s.commit()
    assert c.get("/api/status").status_code == 401


def test_any_google_account_gets_in_without_invitation():
    c = client("cualquiera@example.com")  # nadie lo invitó ni lo conoce el dueño
    assert c.get("/api/status").status_code == 200


def test_logout_closes_session():
    c = client(OWNER)
    assert c.get("/api/status").status_code == 200
    assert c.post("/api/auth/logout").status_code == 200
    with Session(engine) as s:
        assert s.exec(select(LoginSession)).all() == []
    assert c.get("/api/status").status_code == 401


# ---------------------------------------------------------------- flujo de la app (con el dueño)
def test_flow():
    c = client(OWNER)
    assert c.get("/api/status").json()["usuario"]["es_dueno"] is True
    # 8 h de jornada; día con 9h30 → 1h30 extra
    r = c.post("/api/shifts", json={"fecha": "2026-09-21", "inicio": "08:00", "fin": "17:30", "nota": "a"})
    assert r.status_code == 200, r.text
    assert r.json()["duracion"] == "09:30:00"
    assert r.json()["inicio"].endswith("-03:00")
    # dos tramos el mismo día: 5h + 4h = 9h → 1h extra
    c.post("/api/shifts", json={"fecha": "2026-09-22", "inicio": "08:00", "fin": "13:00"})
    sid = c.post("/api/shifts", json={"fecha": "2026-09-22", "inicio": "14:00", "fin": "18:00"}).json()["id"]
    # turno nocturno que cruza medianoche: 22:00 → 06:00 = 8h, 0 extra
    c.post("/api/shifts", json={"fecha": "2026-09-23", "inicio": "22:00", "fin": "06:00"})
    s = c.get("/api/summary").json()
    days = {d["fecha"]: d for d in s["dias"]}
    assert days["2026-09-21"]["extra"] == "01:30:00"
    assert days["2026-09-22"]["extra"] == "01:00:00"
    assert days["2026-09-23"]["trabajado"] == "08:00:00"
    b = c.get("/api/bank").json()
    assert b["saldo"] == "02:30:00"
    # modificar: salida 19:00 → segundo tramo 5h → 2h extra ese día
    c.put(f"/api/shifts/{sid}", json={"fin": "19:00"})
    assert c.get("/api/bank").json()["saldo"] == "03:30:00"
    # corrección manual de extras del 21 a 0h45
    c.put("/api/days/2026-09-21/extra", json={"horas": 0, "minutos": 45})
    assert c.get("/api/bank").json()["saldo"] == "02:45:00"
    # usar 2 h del banco
    m = c.post("/api/bank", json={"fecha": "2026-09-25", "horas": 2, "nota": "trámite"}).json()
    assert m["segundos"] == -7200
    b = c.get("/api/bank").json()
    assert b["saldo"] == "00:45:00"
    assert b["desglose"]["laboral"]["minutos"] == 45
    # usar 1 día (8 h) → saldo negativo
    c.post("/api/bank", json={"fecha": "2026-09-26", "dias": 1})
    b = c.get("/api/bank").json()
    assert b["saldo"] == "-07:15:00" and b["desglose"]["laboral"]["signo"] == -1
    # fichar entrada/salida
    assert c.post("/api/clock-in", json={"hora": "2026-09-24T09:00"}).status_code == 200
    assert c.post("/api/clock-in", json={}).status_code == 400
    r = c.post("/api/clock-out", json={"hora": "12:00"})
    assert r.status_code == 200 and r.json()["duracion"] == "03:00:00"
    # excel
    x = c.get("/api/export.xlsx")
    assert x.status_code == 200
    wb = load_workbook(BytesIO(x.content))
    assert wb.sheetnames == ["Jornadas", "Resumen diario", "Banco de horas"]
    # ajustes
    assert c.put("/api/settings", json={"daily_hours": 7}).json()["daily_hours"] == 7
    assert c.put("/api/settings", json={"timezone": "Nope/X"}).status_code == 400


# ---------------------------------------------------------------- cada usuario ve solo lo suyo
def test_data_is_isolated_between_users():
    boss, friend = client(OWNER), client(FRIEND)
    sid = boss.post("/api/shifts", json={"fecha": "2026-09-21", "inicio": "08:00", "fin": "17:30"}).json()["id"]
    mid = boss.post("/api/bank", json={"fecha": "2026-09-25", "horas": 1}).json()["id"]
    boss.put("/api/days/2026-09-21/extra", json={"horas": 3})
    boss.put("/api/settings", json={"daily_hours": 6})

    # el amigo no ve nada de lo del dueño
    assert friend.get("/api/shifts").json() == []
    assert friend.get("/api/summary").json()["dias"] == []
    assert friend.get("/api/bank").json()["saldo"] == "00:00:00"
    assert friend.get("/api/settings").json()["daily_hours"] == 8  # sus ajustes por defecto, no los del dueño
    # ni puede tocarlo aunque conozca los ids
    assert friend.put(f"/api/shifts/{sid}", json={"fin": "23:00"}).status_code == 400
    assert friend.delete(f"/api/shifts/{sid}").status_code == 400
    assert friend.put(f"/api/bank/{mid}", json={"horas": 9}).status_code == 400
    assert friend.delete(f"/api/bank/{mid}").status_code == 400
    # lo del dueño sigue intacto
    assert boss.get("/api/shifts").json()[0]["duracion"] == "09:30:00"
    assert boss.get("/api/bank").json()["movimientos"][0]["segundos"] == -3600

    # el mismo día puede tener corrección para cada uno
    friend.put("/api/days/2026-09-21/extra", json={"horas": 1})
    assert [d["extra"] for d in boss.get("/api/summary").json()["dias"]] == ["03:00:00"]
    assert [d["extra"] for d in friend.get("/api/summary").json()["dias"]] == ["01:00:00"]
    # cada uno puede tener su propia jornada abierta
    assert boss.post("/api/clock-in", json={"hora": "2026-09-24T09:00"}).status_code == 200
    assert friend.post("/api/clock-in", json={"hora": "2026-09-24T10:00"}).status_code == 200
    # el Excel de cada uno solo trae lo suyo
    rows = lambda c: load_workbook(BytesIO(c.get("/api/export.xlsx").content))["Jornadas"].max_row  # noqa: E731
    assert rows(boss) == 4 and rows(friend) == 3  # encabezado + jornadas + fila de total


# ---------------------------------------------------------------- acceso y bloqueo
def test_owner_manages_access():
    boss, friend = client(OWNER), client(FRIEND)
    friend.post("/api/shifts", json={"fecha": "2026-09-21", "inicio": "08:00", "fin": "09:00"})  # ya entró y cargó algo
    assert [u["email"] for u in boss.get("/api/admin/usuarios").json()] == [FRIEND]
    assert boss.post("/api/admin/bloqueados", json={"email": "no-es-correo"}).status_code == 400
    assert boss.post("/api/admin/bloqueados", json={"email": OWNER}).status_code == 400  # no se puede bloquear a sí mismo
    # solo el dueño administra
    assert friend.get("/api/admin/usuarios").status_code == 403
    assert friend.post("/api/admin/bloqueados", json={"email": "otro@example.com"}).status_code == 403
    assert friend.get("/api/status").json()["usuario"]["es_dueno"] is False
    # bloquear corta el acceso al instante (se cierra la sesión), pero sus datos se conservan
    r = boss.post("/api/admin/bloqueados", json={"email": "  Amigo@Example.com "})
    assert r.status_code == 200 and [(u["email"], u["bloqueado"]) for u in r.json()] == [(FRIEND, True)]
    assert friend.get("/api/status").status_code == 401
    # desbloquear le permite volver a entrar (con una sesión nueva)
    assert [u["bloqueado"] for u in boss.delete(f"/api/admin/bloqueados/{FRIEND}").json()] == [False]
    assert len(client(FRIEND).get("/api/shifts").json()) == 1
    # se puede bloquear a alguien que todavía nunca entró
    r = boss.post("/api/admin/bloqueados", json={"email": "nunca-entro@example.com"})
    nuevo = next(u for u in r.json() if u["email"] == "nunca-entro@example.com")
    assert nuevo["bloqueado"] is True and nuevo["nombre"] == "" and nuevo["ultimo_ingreso"] is None
    assert client("nunca-entro@example.com").get("/api/status").status_code == 401


# ---------------------------------------------------------------- ingreso con Google
class FakeCreds:
    id_token = "id-token-falso"

    def to_json(self):
        return json.dumps({"token": "t", "refresh_token": "r", "client_id": "c", "client_secret": "s",
                           "token_uri": "https://oauth2.googleapis.com/token"})


class FakeFlow:
    code_verifier = None
    credentials = FakeCreds()

    def fetch_token(self, code):
        assert code == "codigo"


def google_setup(monkeypatch, email: str):
    monkeypatch.setenv("GOOGLE_CLIENT_ID", "id")
    monkeypatch.setenv("GOOGLE_CLIENT_SECRET", "secreto")
    monkeypatch.setenv("GOOGLE_REDIRECT_URI", "http://testserver/api/google/callback")
    monkeypatch.setattr(gi, "_verify_id_token",
                        lambda t: {"email": email, "email_verified": True, "name": "Ana", "picture": "http://x/a.png"})


def start_login(c: TestClient) -> str:
    r = c.get("/api/auth/google/start", follow_redirects=False)
    assert r.status_code in (302, 307) and r.headers["location"].startswith("https://accounts.google.com/")
    with Session(engine) as s:
        return s.exec(select(OAuthState)).one().state


def test_google_login_creates_session(monkeypatch):
    google_setup(monkeypatch, OWNER)
    c = client()
    state = start_login(c)
    monkeypatch.setattr(gi, "_flow", lambda state=None: FakeFlow())
    r = c.get("/api/google/callback", params={"code": "codigo", "state": state}, follow_redirects=False)
    assert r.headers["location"].endswith("?login=ok")
    assert auth.COOKIE in r.headers["set-cookie"] and "HttpOnly" in r.headers["set-cookie"]
    me = c.get("/api/status").json()
    assert me["usuario"]["email"] == OWNER and me["usuario"]["nombre"] == "Ana" and me["google"]["conectado"] is True
    with Session(engine) as s:
        assert s.exec(select(OAuthState)).all() == []  # el estado se usa una sola vez
        assert json.loads(s.exec(select(GoogleToken)).one().credentials_json)["refresh_token"] == "r"
    # el estado usado no sirve de nuevo
    r = c.get("/api/google/callback", params={"code": "codigo", "state": state}, follow_redirects=False)
    assert r.headers["location"].endswith("?login=error")


def test_google_login_denied_for_blocked_account(monkeypatch):
    block("bloqueado@example.com")
    google_setup(monkeypatch, "bloqueado@example.com")
    c = client()
    state = start_login(c)
    monkeypatch.setattr(gi, "_flow", lambda state=None: FakeFlow())
    r = c.get("/api/google/callback", params={"code": "codigo", "state": state}, follow_redirects=False)
    assert r.headers["location"].endswith("?login=denied") and "set-cookie" not in r.headers
    with Session(engine) as s:  # no se crea el usuario ni se guardan sus permisos
        assert s.exec(select(User)).all() == [] and s.exec(select(GoogleToken)).all() == []
    assert c.get("/api/status").status_code == 401


def test_google_login_open_access_and_bad_state(monkeypatch):
    """Cualquier cuenta de Google entra, sin que nadie la haya invitado."""
    google_setup(monkeypatch, FRIEND)
    c = client()
    monkeypatch.setattr(gi, "_flow", lambda state=None: FakeFlow())
    r = c.get("/api/google/callback", params={"code": "codigo", "state": "inventado"}, follow_redirects=False)
    assert r.headers["location"].endswith("?login=error")
    monkeypatch.undo()
    google_setup(monkeypatch, FRIEND)
    state = start_login(c)
    monkeypatch.setattr(gi, "_flow", lambda state=None: FakeFlow())
    r = c.get("/api/google/callback", params={"code": "codigo", "state": state}, follow_redirects=False)
    assert r.headers["location"].endswith("?login=ok")
    assert c.get("/api/status").json()["usuario"]["email"] == FRIEND


def test_google_start_without_config_redirects_with_notice(monkeypatch):
    monkeypatch.setenv("GOOGLE_CLIENT_ID", "")
    r = client().get("/api/auth/google/start", follow_redirects=False)
    assert r.headers["location"].endswith("?login=config")


def test_google_disconnect_only_affects_the_user(monkeypatch):
    for email in (OWNER, FRIEND):
        uid, _ = make_user(email)
        with Session(engine) as s:
            s.add(GoogleToken(user_id=uid, credentials_json='{"token": "t"}'))
            s.commit()
    monkeypatch.setattr("requests.post", lambda *a, **k: None)  # no llama a Google
    boss, friend = client(OWNER), client(FRIEND)
    assert boss.post("/api/google/disconnect").status_code == 200
    assert boss.get("/api/status").json()["google"]["conectado"] is False
    assert friend.get("/api/status").json()["google"]["conectado"] is True


# ---------------------------------------------------------------- alarmas
def google_connect(uid: int) -> None:
    with Session(engine) as s:
        s.add(GoogleToken(user_id=uid, credentials_json='{"token": "t"}'))
        s.commit()


def test_alarms_crud_and_isolation():
    boss, friend = client(OWNER), client(FRIEND)
    r = boss.post("/api/alarms", json={"nombre": "Despertar", "cuando": "07:00", "motivo": "Arrancar el día",
                                       "sonido": "urgente", "veces": 3, "intervalo_segundos": 10})
    assert r.status_code == 200, r.text
    a = r.json()
    assert (a["nombre"], a["motivo"], a["sonido"], a["veces"], a["intervalo_segundos"]) == ("Despertar", "Arrancar el día", "urgente", 3, 10)
    assert a["hora"].endswith("07:00:00-03:00") and a["en_calendar"] is False and a["confirmada"] is False
    aid = a["id"]
    assert [x["id"] for x in boss.get("/api/alarms").json()] == [aid]
    # validaciones
    assert boss.post("/api/alarms", json={"nombre": "", "cuando": "07:00"}).status_code == 400
    assert boss.post("/api/alarms", json={"nombre": "x", "cuando": "07:00", "sonido": "rara"}).status_code == 400
    assert boss.post("/api/alarms", json={"nombre": "x", "cuando": "07:00", "veces": 0}).status_code == 400
    assert boss.post("/api/alarms", json={"nombre": "x", "cuando": "07:00", "veces": 999}).status_code == 400
    assert boss.post("/api/alarms", json={"nombre": "x", "cuando": "07:00", "intervalo_segundos": 1}).status_code == 400
    # modificar
    r = boss.put(f"/api/alarms/{aid}", json={"cuando": "08:30", "veces": 5})
    assert r.status_code == 200 and r.json()["hora"].endswith("08:30:00-03:00") and r.json()["veces"] == 5
    # otro usuario no la ve ni puede tocarla
    assert friend.get("/api/alarms").json() == []
    assert friend.put(f"/api/alarms/{aid}", json={"veces": 1}).status_code == 400
    assert friend.delete(f"/api/alarms/{aid}").status_code == 400
    assert boss.delete(f"/api/alarms/{aid}").status_code == 200
    assert boss.get("/api/alarms").json() == []


def test_alarm_dismiss_confirms_and_resets_on_reschedule():
    boss = client(OWNER)
    aid = boss.post("/api/alarms", json={"nombre": "Reunión", "cuando": "09:00"}).json()["id"]
    assert boss.post(f"/api/alarms/{aid}/dismiss").json()["confirmada"] is True
    # reprogramarla la vuelve a dejar pendiente
    assert boss.put(f"/api/alarms/{aid}", json={"cuando": "10:00"}).json()["confirmada"] is False


def test_alarm_reinforces_with_calendar_event_and_stays_in_sync(monkeypatch):
    uid, _ = make_user(OWNER)
    boss = client(OWNER)
    events = []
    monkeypatch.setattr(gi, "create_event", lambda s, u, titulo, inicio, **k: (events.append(("create", titulo)), {"id": "ev1", "htmlLink": "x"})[1])
    monkeypatch.setattr(gi, "update_event", lambda s, u, event_id, **k: events.append(("update", event_id, k.get("titulo"))))
    monkeypatch.setattr(gi, "delete_event", lambda s, u, event_id: events.append(("delete", event_id)))

    # sin Google conectado: no se intenta nada, y la alarma igual se crea bien
    r = boss.post("/api/alarms", json={"nombre": "Sin Google", "cuando": "07:00"})
    assert r.status_code == 200 and r.json()["en_calendar"] is False and events == []
    boss.delete(f"/api/alarms/{r.json()['id']}")

    google_connect(uid)
    r = boss.post("/api/alarms", json={"nombre": "Despertar", "cuando": "07:00"})
    aid = r.json()["id"]
    assert r.json()["en_calendar"] is True and events == [("create", "⏰ Despertar")]
    boss.put(f"/api/alarms/{aid}", json={"cuando": "08:00"})
    assert events[-1] == ("update", "ev1", "⏰ Despertar")
    boss.delete(f"/api/alarms/{aid}")
    assert events[-1] == ("delete", "ev1")


def test_alarm_calendar_failure_does_not_break_the_alarm(monkeypatch):
    uid, _ = make_user(OWNER)
    google_connect(uid)
    monkeypatch.setattr(gi, "create_event", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("Google caído")))
    r = client(OWNER).post("/api/alarms", json={"nombre": "Igual funciona", "cuando": "07:00"})
    assert r.status_code == 200 and r.json()["en_calendar"] is False


def test_assistant_alarm_tools():
    uid, _ = make_user(OWNER)
    other, _ = make_user(FRIEND)
    with Session(engine) as s:
        u, o = s.get(User, uid), s.get(User, other)
        r = assistant.run_tool(s, u, "crear_alarma", {"nombre": "Despertar", "cuando": "07:00", "veces": 4})
        assert r["nombre"] == "Despertar" and r["veces"] == 4
        assert [x["nombre"] for x in assistant.run_tool(s, u, "listar_alarmas", {})] == ["Despertar"]
        assert assistant.run_tool(s, o, "listar_alarmas", {}) == []  # otro usuario no la ve
        m = assistant.run_tool(s, u, "modificar_alarma", {"id": r["id"], "veces": 8})
        assert m["veces"] == 8
        with pytest.raises(ValueError):  # otro usuario no puede tocarla
            assistant.run_tool(s, o, "modificar_alarma", {"id": r["id"], "veces": 1})
        assert "error" in assistant.run_tool(s, u, "eliminar_alarma", {"id": r["id"], "confirmado": False})
        assert assistant.run_tool(s, u, "eliminar_alarma", {"id": r["id"], "confirmado": True}) == {"ok": True}


# ---------------------------------------------------------------- conversación con el asistente
def fake_assistant(monkeypatch):
    """Reemplaza a Gemini: responde con un eco y registra qué historial le llegó."""
    seen = []

    def fake_chat(s, u, message, history=None):
        seen.append((u.email, message, list(history or [])))
        return {"respuesta": f"eco: {message}", "acciones": [{"herramienta": "consultar_banco", "argumentos": {}, "ok": True}]}

    monkeypatch.setattr(assistant, "chat", fake_chat)
    return seen


def test_chat_history_is_the_same_on_every_device(monkeypatch):
    seen = fake_assistant(monkeypatch)
    pc, phone, friend = client(OWNER), client(OWNER), client(FRIEND)  # pc y teléfono: dos sesiones del mismo usuario
    assert phone.get("/api/assistant/history").json() == []
    assert pc.post("/api/assistant", json={"mensaje": "hola"}).json()["respuesta"] == "eco: hola"
    h = phone.get("/api/assistant/history").json()  # lo escrito en la PC aparece en el teléfono
    assert [(m["role"], m["text"]) for m in h] == [("user", "hola"), ("assistant", "eco: hola")]
    assert h[1]["acciones"][0]["herramienta"] == "consultar_banco" and h[0]["acciones"] == []
    # el modelo recibe la conversación previa desde el servidor, no desde el navegador
    phone.post("/api/assistant", json={"mensaje": "segunda", "historial": [{"role": "user", "text": "inventado"}]})
    assert seen[1][2] == [{"role": "user", "text": "hola"}, {"role": "assistant", "text": "eco: hola"}]
    # cada persona ve solo la suya
    assert friend.get("/api/assistant/history").json() == []
    friend.post("/api/assistant", json={"mensaje": "soy el amigo"})
    assert seen[2][2] == []
    assert [m["text"] for m in pc.get("/api/assistant/history").json()] == ["hola", "eco: hola", "segunda", "eco: segunda"]
    # borrar la conversación no toca la de otro
    assert friend.delete("/api/assistant/history").json() == {"ok": True}
    assert friend.get("/api/assistant/history").json() == []
    assert len(pc.get("/api/assistant/history").json()) == 4
    assert pc.delete("/api/assistant/history").status_code == 200
    assert phone.get("/api/assistant/history").json() == []


def test_chat_history_requires_login_and_limits_size(monkeypatch):
    fake_assistant(monkeypatch)
    anon = client()
    assert anon.get("/api/assistant/history").status_code == 401
    assert anon.delete("/api/assistant/history").status_code == 401
    boss = client(OWNER)
    assert boss.post("/api/assistant", json={"mensaje": "x" * 2001}).status_code == 422
    assert boss.post("/api/assistant", json={"mensaje": ""}).status_code == 422
    assert boss.get("/api/assistant/history").json() == []  # los rechazados no se guardan


def test_chat_history_keeps_only_the_latest_messages():
    uid, _ = make_user(OWNER)
    with Session(engine) as s:
        for i in range(1, history.MAX_MESSAGES // 2 + 11):  # 10 intercambios más de los que entran
            history.add_exchange(s, uid, f"msg {i}", f"resp {i}", [])
        rows = history.recent(s, uid)
        assert len(rows) == history.MAX_MESSAGES
        assert rows[0].text == "msg 11" and rows[-1].text == f"resp {history.MAX_MESSAGES // 2 + 10}"
        assert len(history.context(s, uid)) == history.CONTEXT


# ---------------------------------------------------------------- uso de la IA y límites de Google
def fake_gemini(monkeypatch, fail=None):
    """Reemplaza al cliente de Gemini. Sin fail: pide una herramienta y luego responde (2 solicitudes de 100+20 tokens)."""
    from google import genai
    from google.genai import types
    meta = types.GenerateContentResponseUsageMetadata(prompt_token_count=100, candidates_token_count=20, total_token_count=120)
    calls = []

    class FakeModels:
        def generate_content(self, model, contents, config):
            calls.append(1)
            if fail:
                raise fail
            part = (types.Part(function_call=types.FunctionCall(name="consultar_banco", args={})) if len(calls) % 2
                    else types.Part(text="Listo."))
            return types.GenerateContentResponse(candidates=[types.Candidate(content=types.Content(role="model", parts=[part]))],
                                                 usage_metadata=meta)

    class FakeClient:
        def __init__(self, api_key):
            self.models = FakeModels()

    monkeypatch.setenv("GEMINI_API_KEY", "x")
    monkeypatch.setattr(genai, "Client", FakeClient)


def test_ai_usage_is_counted_per_user(monkeypatch):
    fake_gemini(monkeypatch)
    boss, friend = client(OWNER), client(FRIEND)
    assert client().get("/api/ai-usage").status_code == 401
    assert boss.get("/api/ai-usage").json()["yo"] == {"solicitudes": 0, "tokens_entrada": 0, "tokens_salida": 0}
    assert boss.post("/api/assistant", json={"mensaje": "uno"}).json()["respuesta"] == "Listo."  # 2 solicitudes a Gemini
    boss.post("/api/assistant", json={"mensaje": "dos"})
    friend.post("/api/assistant", json={"mensaje": "tres"})
    mine = boss.get("/api/ai-usage").json()
    assert mine["yo"] == {"solicitudes": 4, "tokens_entrada": 400, "tokens_salida": 80}
    assert mine["dia"] == usage.quota_day() and mine["reinicia"][:2] == "20"
    # el dueño ve el de todos y el total; cada quien más ve solo lo suyo
    assert {x["email"]: x["solicitudes"] for x in mine["todos"]} == {OWNER: 4, FRIEND: 2}
    assert mine["total"] == {"solicitudes": 6, "tokens_entrada": 600, "tokens_salida": 120}
    theirs = friend.get("/api/ai-usage").json()
    assert theirs["yo"]["solicitudes"] == 2 and "todos" not in theirs and "total" not in theirs


def test_assistant_actions_include_the_tool_result(monkeypatch):
    """El resultado completo de cada acción viaja en la respuesta (lo usa, p. ej., el atajo al Reloj de Android)."""
    from google import genai
    from google.genai import types
    meta = types.GenerateContentResponseUsageMetadata(prompt_token_count=10, candidates_token_count=5, total_token_count=15)

    class FakeModels:
        def generate_content(self, model, contents, config):
            if len(contents) == 1:
                fc = types.FunctionCall(name="crear_alarma", args={"nombre": "Despertar", "cuando": "07:00"})
                return types.GenerateContentResponse(candidates=[types.Candidate(content=types.Content(
                    role="model", parts=[types.Part(function_call=fc)]))], usage_metadata=meta)
            return types.GenerateContentResponse(candidates=[types.Candidate(content=types.Content(
                role="model", parts=[types.Part(text="Lista.")]))], usage_metadata=meta)

    class FakeClient:
        def __init__(self, api_key):
            self.models = FakeModels()

    monkeypatch.setenv("GEMINI_API_KEY", "x")
    monkeypatch.setattr(genai, "Client", FakeClient)
    r = client(OWNER).post("/api/assistant", json={"mensaje": "poneme una alarma a las 7"})
    assert r.status_code == 200
    accion = r.json()["acciones"][0]
    assert accion["herramienta"] == "crear_alarma" and accion["ok"] is True
    assert accion["resultado"]["nombre"] == "Despertar" and accion["resultado"]["hora"].endswith("07:00:00-03:00")


def test_quota_errors_become_clear_messages(monkeypatch):
    from google.genai import errors
    boss = client(OWNER)
    fake_gemini(monkeypatch, fail=errors.APIError(429, {"error": {"code": 429, "status": "RESOURCE_EXHAUSTED", "message": "quota"}}))
    r = boss.post("/api/assistant", json={"mensaje": "hola"})
    assert r.status_code == 429 and "límite gratuito" in r.json()["detail"] and "se reinicia a las" in r.json()["detail"]
    assert "RESOURCE_EXHAUSTED" not in r.json()["detail"]
    fake_gemini(monkeypatch, fail=errors.APIError(503, {"error": {"code": 503, "status": "UNAVAILABLE", "message": "busy"}}))
    r = boss.post("/api/assistant", json={"mensaje": "hola"})
    assert r.status_code == 503 and "mucha demanda" in r.json()["detail"]
    fake_gemini(monkeypatch, fail=errors.APIError(500, {"error": {"code": 500, "message": "boom"}}))
    r = boss.post("/api/assistant", json={"mensaje": "hola"})
    assert r.status_code == 502 and r.json()["detail"].startswith("Error del asistente")
    # lo que falló no cuenta como uso ni queda en la conversación
    assert boss.get("/api/ai-usage").json()["yo"]["solicitudes"] == 0
    assert boss.get("/api/assistant/history").json() == []


def test_quota_day_follows_pacific_midnight():
    from datetime import datetime, timezone
    from zoneinfo import ZoneInfo
    utc = timezone.utc
    assert usage.quota_day(datetime(2026, 9, 27, 6, 59, tzinfo=utc)) == "2026-09-26"  # todavía 23:59 en el Pacífico
    assert usage.quota_day(datetime(2026, 9, 27, 7, 0, tzinfo=utc)) == "2026-09-27"
    ar = ZoneInfo("America/Argentina/San_Juan")
    verano = usage.next_reset(ar, datetime(2026, 9, 27, 12, 0, tzinfo=utc))  # EE. UU. con horario de verano
    invierno = usage.next_reset(ar, datetime(2026, 12, 1, 12, 0, tzinfo=utc))
    assert (verano.hour, verano.day) == (4, 28) and (invierno.hour, invierno.day) == (5, 2)


# ---------------------------------------------------------------- lógica
def test_parse_time_without_leading_zero():
    from datetime import date

    class St:
        timezone = "America/Argentina/San_Juan"  # UTC-3
    for txt in ("7:45", "07:45", "7:45:00"):
        assert calc.parse_dt(txt, St, date(2026, 9, 25)).isoformat() == "2026-09-25T10:45:00"


def test_dependencies_in_sync():
    """Vercel instala desde pyproject.toml y en local se usa requirements.txt: deben listar lo mismo."""
    import tomllib
    from pathlib import Path
    root = Path(__file__).resolve().parents[2]
    in_pyproject = tomllib.loads((root / "pyproject.toml").read_text(encoding="utf-8"))["project"]["dependencies"]
    in_requirements = [ln.strip() for ln in (root / "requirements.txt").read_text(encoding="utf-8").splitlines()
                       if ln.strip() and not ln.startswith("#")]
    assert sorted(in_pyproject) == sorted(in_requirements)


def test_breakdown():
    class St:  # jornada 8 h, mes 22 jornadas
        daily_hours, workdays_per_month = 8, 22
    secs = 23 * 8 * 3600 + 3 * 3600 + 5 * 60 + 7  # 23 jornadas + 3h 5m 7s
    lab = calc.bank_breakdowns(secs, St)["laboral"]
    assert (lab["meses"], lab["dias"], lab["horas"], lab["minutos"], lab["segundos"]) == (1, 1, 3, 5, 7)
    rel = calc.bank_breakdowns(90061, St)["reloj"]  # 1 día 1 h 1 min 1 s
    assert (rel["meses"], rel["dias"], rel["horas"], rel["minutos"], rel["segundos"]) == (0, 1, 1, 1, 1)


# ---------------------------------------------------------------- asistente
def test_assistant_tools():
    from google.genai import types
    from backend.app import assistant
    for t in assistant.TOOLS:
        types.FunctionDeclaration(**t)
    uid, _ = make_user(OWNER, "Ana")
    other, _ = make_user(FRIEND)
    with Session(engine) as s:
        u, o = s.get(User, uid), s.get(User, other)
        r = assistant.run_tool(s, u, "registrar_jornada", {"fecha": "2026-08-03", "inicio": "08:00", "fin": "18:00"})
        assert r["duracion"] == "10:00:00"
        assistant.run_tool(s, u, "fijar_extras_dia", {"fecha": "2026-08-03", "horas": 1, "minutos": 30})
        d = assistant.run_tool(s, u, "resumen", {"desde": "2026-08-03", "hasta": "2026-08-03"})
        assert d["dias"][0]["extra"] == "01:30:00"
        # otro usuario no ve ni puede borrar esa jornada
        assert assistant.run_tool(s, o, "listar_jornadas", {}) == []
        with pytest.raises(ValueError):
            assistant.run_tool(s, o, "eliminar_jornada", {"id": r["id"], "confirmado": True})
        assert "error" in assistant.run_tool(s, u, "eliminar_jornada", {"id": r["id"], "confirmado": False})
        assert assistant.run_tool(s, u, "eliminar_jornada", {"id": r["id"], "confirmado": True}) == {"ok": True}
        assert "desglose" in assistant.run_tool(s, u, "consultar_banco", {})
        assert assistant._system_prompt(s, u).startswith("Sos el asistente") and "Ana" in assistant._system_prompt(s, u)


def test_chat_loop(monkeypatch):
    """Simula a Gemini: primero pide una herramienta y luego responde texto."""
    from google import genai
    from google.genai import types
    from backend.app import assistant

    calls = []

    class FakeModels:
        def generate_content(self, model, contents, config):
            calls.append(contents)
            if len(calls) == 1:
                return types.GenerateContentResponse(candidates=[types.Candidate(content=types.Content(
                    role="model", parts=[types.Part(function_call=types.FunctionCall(name="consultar_banco", args={}))]))])
            assert contents[-1].parts[0].function_response.name == "consultar_banco"
            return types.GenerateContentResponse(candidates=[types.Candidate(content=types.Content(
                role="model", parts=[types.Part(text="Tenés saldo.")]))])

    class FakeClient:
        def __init__(self, api_key):
            self.models = FakeModels()

    monkeypatch.setenv("GEMINI_API_KEY", "x")
    monkeypatch.setattr(genai, "Client", FakeClient)
    uid, _ = make_user(OWNER)
    with Session(engine) as s:
        out = assistant.chat(s, s.get(User, uid), "¿cuánto tengo en el banco?", [{"role": "assistant", "text": "hola"}])
    assert out["respuesta"] == "Tenés saldo." and out["acciones"][0]["ok"]
